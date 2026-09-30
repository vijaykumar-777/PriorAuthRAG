#!/usr/bin/env python3
"""
scripts/compare_test_results.py
--------------------------------------------------------------------------------
Phase 5.4: Automated Comparison Engine & Reporting Suite

Reconciles the 50 Prior Authorization batch test execution results
(data/test_runs/latest/system_run_results.json) against the Ground Truth
specifications (data/synthetic_cases/Synthetic_Test_Cases.xlsx and
data/ground_truth/Policy_Ground_Truth.xlsx).

Outputs:
1. data/results/PA_Test_Results.xlsx (Multi-sheet styled workbook):
   - Case_Level_Results (50 rows)
   - Criterion_Level_Results (170+ rows)
   - Summary_Metrics (Executive KPI scorecard)
2. data/results/metrics_summary.json (Machine-readable metrics)
--------------------------------------------------------------------------------
"""

import os
import sys
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WORKSPACE_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SYSTEM_RUNS_PATH = os.path.join(WORKSPACE_ROOT, "data", "test_runs", "latest", "system_run_results.json")
SYNTHETIC_CASES_PATH = os.path.join(WORKSPACE_ROOT, "data", "synthetic_cases", "Synthetic_Test_Cases.xlsx")
POLICY_GT_PATH = os.path.join(WORKSPACE_ROOT, "data", "ground_truth", "Policy_Ground_Truth.xlsx")
OUTPUT_EXCEL_PATH = os.path.join(WORKSPACE_ROOT, "data", "results", "PA_Test_Results.xlsx")
OUTPUT_METRICS_PATH = os.path.join(WORKSPACE_ROOT, "data", "results", "metrics_summary.json")

def load_inputs():
    print(f"Loading system run results from: {SYSTEM_RUNS_PATH}")
    with open(SYSTEM_RUNS_PATH, "r", encoding="utf-8") as f:
        sys_data = json.load(f)
    sys_runs = {c["case_id"]: c for c in sys_data}

    print(f"Loading synthetic test cases from: {SYNTHETIC_CASES_PATH}")
    df_cases_gt = pd.read_excel(SYNTHETIC_CASES_PATH, sheet_name="Case_Level")
    df_crit_gt = pd.read_excel(SYNTHETIC_CASES_PATH, sheet_name="Criterion_Level")

    print(f"Loading policy ground truth from: {POLICY_GT_PATH}")
    df_policy_gt = pd.read_excel(POLICY_GT_PATH, sheet_name="Policy_Ground_Truth")

    return sys_runs, df_cases_gt, df_crit_gt, df_policy_gt

def normalize_status(status_str):
    if not status_str:
        return "UNKNOWN"
    s = str(status_str).strip().upper()
    if "CONFLICT" in s:
        return "CONFLICT"
    if "NOT MET" in s or "DENIED" in s or "NOT_MET" in s:
        return "NOT MET"
    if "MET" in s or "APPROVED" in s:
        return "MET"
    if "UNKNOWN" in s or "MISSING" in s:
        return "UNKNOWN"
    return s

def evaluate_case_match(expected, actual, archetype):
    exp_norm = normalize_status(expected)
    act_norm = normalize_status(actual)

    if exp_norm == act_norm:
        return True, "PASS", "Exact match with ground truth"
    
    # Archetype 8: Conflicting Evidence
    if exp_norm == "CONFLICT" and act_norm == "CONFLICT":
        return True, "PASS", "Correctly identified conflicting clinical evidence"
    
    # Archetype 3: Information Missing
    if exp_norm == "UNKNOWN" and act_norm in ["UNKNOWN", "NOT MET"]:
        return True, "PASS", f"Missing information recognized (System yielded {act_norm})"

    # Discrepancies
    if exp_norm == "MET" and act_norm != "MET":
        return False, "FAIL", f"Improper Denial (False Negative): Expected MET, got {act_norm}"
    elif exp_norm != "MET" and act_norm == "MET":
        return False, "FAIL", f"Inappropriate Approval (False Positive): Expected {exp_norm}, got MET"
    else:
        return False, "FAIL", f"Determination discrepancy: Expected {exp_norm}, got {act_norm}"

def classify_criterion_error(gt_res, sys_res, archetype, crit_desc):
    if sys_res == "NOT_EVALUATED":
        return "RETRIEVAL_ERROR", "Criterion was not retrieved or evaluated by the RAG policy engine"
    
    gt_norm = normalize_status(gt_res)
    sys_norm = normalize_status(sys_res)

    if gt_norm == sys_norm:
        return "NONE", "Matched ground truth"

    if "Conflicting" in str(archetype):
        return "CONFLICT_DETECTION_ERROR", "Failed to detect discordant imaging/physician statements"
    elif "Missing" in str(archetype):
        return "MISSING_EVIDENCE_ERROR", "Failed to recognize absent documentation or required test"
    elif gt_norm == "NOT MET" and sys_norm == "MET":
        return "CLINICAL_REASONING_ERROR", "Overlooked unmet clinical threshold (e.g. conservative therapy duration or imaging severity)"
    elif gt_norm == "MET" and sys_norm == "NOT MET":
        return "CLINICAL_REASONING_ERROR", "False negative clinical synthesis; patient evidence satisfied requirement"
    else:
        return "CLINICAL_REASONING_ERROR", f"Reasoning discrepancy between expected {gt_norm} and system {sys_norm}"

def main():
    print("=" * 80)
    print("STARTING PRIOR AUTHORIZATION COMPARISON ENGINE")
    print("=" * 80)

    sys_runs, df_cases_gt, df_crit_gt, df_policy_gt = load_inputs()

    # -------------------------------------------------------------------------
    # 1. CASE LEVEL COMPARISON
    # -------------------------------------------------------------------------
    case_rows = []
    tp, tn, fp, fn = 0, 0, 0, 0
    passed_cases = 0

    for _, row in df_cases_gt.iterrows():
        cid = str(row["Case ID"]).strip()
        payer = str(row.get("Payer", "Cigna")).strip()
        cpt = str(row.get("CPT", "")).strip()
        cond = str(row.get("Condition", "")).strip()
        archetype = str(row.get("Test Case Type", "")).strip()
        gt_res = str(row.get("Expected Overall Result", "")).strip()

        sys_case = sys_runs.get(cid, {})
        sys_res = sys_case.get("system_overall_result", "NOT_RUN")
        runtime = sys_case.get("runtime_seconds", 0.0)

        is_match, pass_fail, failure_notes = evaluate_case_match(gt_res, sys_res, archetype)

        if is_match:
            passed_cases += 1

        gt_norm = normalize_status(gt_res)
        sys_norm = normalize_status(sys_res)

        if gt_norm == "MET":
            if sys_norm == "MET":
                tp += 1
            else:
                fn += 1
        else: # Ground Truth is NOT MET, CONFLICT, or UNKNOWN
            if sys_norm == "MET":
                fp += 1
            else:
                tn += 1

        case_rows.append({
            "Case ID": cid,
            "Payer": payer,
            "CPT": cpt,
            "Condition": cond,
            "Test Case Type": archetype,
            "Ground Truth Overall Result": gt_res,
            "System Overall Result": sys_res,
            "Overall PASS/FAIL": pass_fail,
            "Runtime (s)": round(runtime, 2),
            "Notes": failure_notes
        })

    # -------------------------------------------------------------------------
    # 2. CRITERION LEVEL COMPARISON
    # -------------------------------------------------------------------------
    crit_rows = []
    passed_crit = 0
    valid_citations = 0

    for _, row in df_crit_gt.iterrows():
        cid = str(row["Case ID"]).strip()
        crit_id = str(row["Criterion ID"]).strip()
        crit_desc = str(row.get("Criterion", "")).strip()
        gt_res = str(row.get("Expected Result", "")).strip()
        pt_ev_expected = str(row.get("Patient Evidence", "")).strip()
        gt_notes = str(row.get("Notes", "")).strip()

        # Find matching evaluated criterion in system run
        sys_case = sys_runs.get(cid, {})
        sys_crits = sys_case.get("criterion_evaluations", [])
        
        # Match by criterion_id or order
        matched_eval = None
        for sc in sys_crits:
            if sc.get("criterion_id") == crit_id:
                matched_eval = sc
                break
        
        # If not matched by explicit ID, fallback to indexed position (e.g. C1 -> idx 0)
        if not matched_eval and crit_id.startswith("C") and crit_id[1:].isdigit():
            idx = int(crit_id[1:]) - 1
            if idx < len(sys_crits):
                matched_eval = sys_crits[idx]

        if matched_eval:
            sys_res = matched_eval.get("system_result", "UNKNOWN")
            sys_ev = matched_eval.get("patient_evidence", "")
            pol_cit = matched_eval.get("policy_citation", "")
            # Determine patient citation source (e.g. physician note, MRI)
            pt_cit = "Patient EHR Records"
            for doc_name in ["Physician Note", "MRI Report", "X-Ray Report", "PT Note", "Lab Report"]:
                if doc_name.lower() in sys_ev.lower() or doc_name.lower() in pt_ev_expected.lower():
                    pt_cit = doc_name
                    break
        else:
            sys_res = "NOT_EVALUATED"
            sys_ev = "None (Criterion was omitted during policy scoping)"
            pol_cit = "Not cited"
            pt_cit = "N/A"

        gt_norm = normalize_status(gt_res)
        sys_norm = normalize_status(sys_res)

        # Check match
        if gt_norm == sys_norm:
            crit_pass_fail = "PASS"
            passed_crit += 1
            err_type = "NONE"
            crit_notes = "Criterion determination matches ground truth"
        else:
            crit_pass_fail = "FAIL"
            # Get test case archetype
            case_info = next((c for c in case_rows if c["Case ID"] == cid), {})
            archetype = case_info.get("Test Case Type", "")
            err_type, crit_notes = classify_criterion_error(gt_res, sys_res, archetype, crit_desc)

        # Citation validation check
        if pol_cit and pol_cit != "Not cited" and ("CMM" in pol_cit or "p." in pol_cit.lower() or "page" in pol_cit.lower() or "section" in pol_cit.lower()):
            valid_citations += 1

        crit_rows.append({
            "Case ID": cid,
            "Criterion ID": crit_id,
            "Ground Truth Criterion Result": gt_res,
            "System Criterion Result": sys_res,
            "Criterion PASS/FAIL": crit_pass_fail,
            "System Evidence": sys_ev,
            "Policy Citation": pol_cit,
            "Patient Citation": pt_cit,
            "Error Type": err_type,
            "Notes": crit_notes
        })

    # -------------------------------------------------------------------------
    # 3. METRICS COMPUTATION
    # -------------------------------------------------------------------------
    total_cases = len(case_rows)
    failed_cases = total_cases - passed_cases
    overall_accuracy = passed_cases / total_cases if total_cases > 0 else 0.0

    sensitivity = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    specificity = tn / (tn + fp) if (tn + fp) > 0 else 0.0
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    f1_score = (2 * precision * sensitivity) / (precision + sensitivity) if (precision + sensitivity) > 0 else 0.0

    total_crit = len(crit_rows)
    crit_accuracy = passed_crit / total_crit if total_crit > 0 else 0.0
    citation_precision = valid_citations / total_crit if total_crit > 0 else 0.0

    avg_runtime = sum(c["Runtime (s)"] for c in case_rows) / total_cases if total_cases > 0 else 0.0

    metrics_dict = {
        "total_test_cases": total_cases,
        "cases_passed": passed_cases,
        "cases_failed": failed_cases,
        "overall_case_accuracy": round(overall_accuracy, 4),
        "true_positives": tp,
        "true_negatives": tn,
        "false_positives": fp,
        "false_negatives": fn,
        "sensitivity_recall": round(sensitivity, 4),
        "specificity": round(specificity, 4),
        "precision": round(precision, 4),
        "f1_score": round(f1_score, 4),
        "total_criteria_evaluated": total_crit,
        "criteria_passed": passed_crit,
        "criterion_level_accuracy": round(crit_accuracy, 4),
        "citation_precision": round(citation_precision, 4),
        "average_runtime_seconds": round(avg_runtime, 2)
    }

    # Save JSON summary
    os.makedirs(os.path.dirname(OUTPUT_METRICS_PATH), exist_ok=True)
    with open(OUTPUT_METRICS_PATH, "w", encoding="utf-8") as f:
        json.dump(metrics_dict, f, indent=2)
    print(f"Saved machine-readable metrics to: {OUTPUT_METRICS_PATH}")

    # -------------------------------------------------------------------------
    # 4. EXCEL WORKBOOK GENERATION (PA_Test_Results.xlsx)
    # -------------------------------------------------------------------------
    wb = Workbook()
    
    # Styles
    navy_header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="375623")
    
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="C65911")

    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Sheet 1: Case_Level_Results
    ws_cases = wb.active
    ws_cases.title = "Case_Level_Results"
    ws_cases.views.sheetView[0].showGridLines = True

    case_headers = [
        "Case ID", "Payer", "CPT", "Condition", "Test Case Type",
        "Ground Truth Overall Result", "System Overall Result", "Overall PASS/FAIL",
        "Runtime (s)", "Notes"
    ]
    ws_cases.append(case_headers)

    for cell in ws_cases[1]:
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, cr in enumerate(case_rows, start=2):
        row_vals = [
            cr["Case ID"], cr["Payer"], cr["CPT"], cr["Condition"], cr["Test Case Type"],
            cr["Ground Truth Overall Result"], cr["System Overall Result"], cr["Overall PASS/FAIL"],
            cr["Runtime (s)"], cr["Notes"]
        ]
        ws_cases.append(row_vals)
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_cases.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if col_idx in [1, 2, 3, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # PASS/FAIL Highlight
            if col_idx == 8:
                if val == "PASS":
                    cell.fill = pass_fill
                    cell.font = pass_font
                else:
                    cell.fill = fail_fill
                    cell.font = fail_font

    ws_cases.freeze_panes = "A2"

    # Sheet 2: Criterion_Level_Results
    ws_crit = wb.create_sheet(title="Criterion_Level_Results")
    ws_crit.views.sheetView[0].showGridLines = True

    crit_headers = [
        "Case ID", "Criterion ID", "Ground Truth Criterion Result",
        "System Criterion Result", "Criterion PASS/FAIL", "System Evidence",
        "Policy Citation", "Patient Citation", "Error Type", "Notes"
    ]
    ws_crit.append(crit_headers)

    for cell in ws_crit[1]:
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, cr in enumerate(crit_rows, start=2):
        row_vals = [
            cr["Case ID"], cr["Criterion ID"], cr["Ground Truth Criterion Result"],
            cr["System Criterion Result"], cr["Criterion PASS/FAIL"], cr["System Evidence"],
            cr["Policy Citation"], cr["Patient Citation"], cr["Error Type"], cr["Notes"]
        ]
        ws_crit.append(row_vals)
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_crit.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if col_idx in [1, 2, 3, 4, 5, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            if col_idx == 5:
                if val == "PASS":
                    cell.fill = pass_fill
                    cell.font = pass_font
                else:
                    cell.fill = fail_fill
                    cell.font = fail_font

    ws_crit.freeze_panes = "A2"

    # Sheet 3: Summary_Metrics
    ws_sum = wb.create_sheet(title="Summary_Metrics")
    ws_sum.views.sheetView[0].showGridLines = True

    title_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    
    ws_sum.merge_cells("B2:E2")
    title_cell = ws_sum["B2"]
    title_cell.value = "PRIOR AUTHORIZATION PIPELINE EVALUATION SCORECARD"
    title_cell.fill = title_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 30

    metrics_rows = [
        ("CORE EVALUATION METRICS", "", ""),
        ("Total Synthetic Test Cases", total_cases, "Complete test cohort across 10 CPT codes"),
        ("Cases Passed (Ground Truth Match)", passed_cases, f"{overall_accuracy*100:.1f}% Overall System Accuracy"),
        ("Cases Failed (Discrepancies)", failed_cases, "Cases routed to root-cause error analysis"),
        ("Overall Case Accuracy", f"{overall_accuracy*100:.2f}%", "Ratio of correct overall PA rulings"),
        ("", "", ""),
        ("CLINICAL ADJUDICATION INTEGRITY", "", ""),
        ("True Positives (Approved when MET)", tp, "Correct approvals"),
        ("True Negatives (Denied when NOT MET)", tn, "Correct denials & identified conflicts"),
        ("False Positives (Inappropriate Approvals)", fp, "Severe Clinical Risk (Over-approvals)"),
        ("False Negatives (Improper Denials)", fn, "Provider Friction (Under-approvals)"),
        ("Sensitivity / True Positive Rate", f"{sensitivity*100:.2f}%", "TP / (TP + FN)"),
        ("Specificity / True Negative Rate", f"{specificity*100:.2f}%", "TN / (TN + FP)"),
        ("Clinical Precision", f"{precision*100:.2f}%", "TP / (TP + FP)"),
        ("Adjudication F1-Score", f"{f1_score*100:.2f}%", "Harmonic mean of precision and sensitivity"),
        ("", "", ""),
        ("CRITERION-LEVEL & CITATION QUALITY", "", ""),
        ("Total Clinical Criteria Evaluated", total_crit, "170 ground-truth medical necessity rules"),
        ("Criteria Passed", passed_crit, f"{crit_accuracy*100:.1f}% criterion-level accuracy"),
        ("Criterion-Level Accuracy", f"{crit_accuracy*100:.2f}%", "Rule-by-rule evaluation fidelity"),
        ("Policy Citation Precision", f"{citation_precision*100:.2f}%", "Valid policy section & page provenance"),
        ("Average Pipeline Runtime per Case", f"{avg_runtime:.2f} s", "End-to-end RAG retrieval + clinical LLM determination")
    ]

    section_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

    curr_row = 4
    for label, val, desc in metrics_rows:
        if not label:
            curr_row += 1
            continue
        if not str(val): # Section Header
            ws_sum.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=5)
            c = ws_sum.cell(row=curr_row, column=2, value=label)
            c.fill = section_header_fill
            c.font = section_font
            c.alignment = Alignment(horizontal="left", vertical="center")
            ws_sum.row_dimensions[curr_row].height = 24
        else:
            c_lbl = ws_sum.cell(row=curr_row, column=2, value=label)
            c_lbl.font = Font(name="Calibri", size=10, bold=True)
            c_lbl.border = thin_border
            
            c_val = ws_sum.cell(row=curr_row, column=3, value=val)
            c_val.font = Font(name="Calibri", size=10, bold=True, color="002060")
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.border = thin_border
            
            c_desc = ws_sum.cell(row=curr_row, column=4, value=desc)
            c_desc.font = Font(name="Calibri", size=10, italic=True)
            c_desc.border = thin_border

            ws_sum.row_dimensions[curr_row].height = 20
        curr_row += 1

    # Auto-adjust column widths
    for ws in [ws_cases, ws_crit, ws_sum]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and len(val_str) < 80:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws_sum.column_dimensions["B"].width = 42
    ws_sum.column_dimensions["C"].width = 16
    ws_sum.column_dimensions["D"].width = 45

    os.makedirs(os.path.dirname(OUTPUT_EXCEL_PATH), exist_ok=True)
    wb.save(OUTPUT_EXCEL_PATH)
    print(f"Saved primary evaluation workbook to: {OUTPUT_EXCEL_PATH}")

    # Console summary scorecard
    print("\n" + "=" * 80)
    print("PRIOR AUTHORIZATION PIPELINE EVALUATION REPORT")
    print("=" * 80)
    print(f"Total Cases Evaluated         : {total_cases}")
    print(f"Cases Passed (Match GT)       : {passed_cases} / {total_cases} ({overall_accuracy*100:.1f}%)")
    print(f"Cases Failed (Discrepancies)  : {failed_cases} / {total_cases} ({(1-overall_accuracy)*100:.1f}%)")
    print("-" * 80)
    print(f"True Positives  (TP)          : {tp}")
    print(f"True Negatives  (TN)          : {tn}")
    print(f"False Positives (FP)          : {fp}  (Inappropriate approvals / Over-approval risk)")
    print(f"False Negatives (FN)          : {fn}  (Improper denials / Provider friction)")
    print(f"Sensitivity (Recall)          : {sensitivity*100:.2f}%")
    print(f"Specificity (Selectivity)     : {specificity*100:.2f}%")
    print(f"Precision                     : {precision*100:.2f}%")
    print(f"F1-Score                      : {f1_score*100:.2f}%")
    print("-" * 80)
    print(f"Total Criteria Evaluated      : {total_crit}")
    print(f"Criteria Passed               : {passed_crit} / {total_crit} ({crit_accuracy*100:.1f}%)")
    print(f"Policy Citation Precision     : {citation_precision*100:.2f}%")
    print(f"Average Runtime per Case      : {avg_runtime:.2f}s")
    print("=" * 80)

if __name__ == "__main__":
    main()
